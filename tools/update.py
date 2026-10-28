"""Get the delivery forms from KoBo and work out each facility's status for this round.

Usage (from the project folder):
    py tools/update.py

Needs private/kobo.json  ({"server": ..., "form": ..., "token": ...}),
or, on GitHub, the secrets KOBO_SERVER, KOBO_FORM and KOBO_TOKEN.

Writes:
    private/<round>/submissions.json   everything KoBo sent, names included (never sent to GitHub)
    data/status.csv, data/status.xlsx  one line per facility, NO names or phone numbers (this is public)
    data/updated.json                  when this last ran
"""
import csv
import json
import math
import os
import sys
import urllib.request
from datetime import datetime, timedelta, timezone

import openpyxl
from openpyxl.styles import Font, PatternFill

from common import DATA, ROOT, load_facilities

LIMIT_M = 500  # further than this from the facility = "location doesn't match"
LAGOS = timezone(timedelta(hours=1))  # Lagos time, so GitHub (which runs on UTC) shows the right hour

CONFIRMED = "Receipt confirmed by facility"
MATCH = "Reported delivered, location matches"
NO_MATCH = "Reported delivered, location doesn't match"
NOT_YET = "Not reported yet"
RANK = {CONFIRMED: 3, MATCH: 2, NO_MATCH: 1}  # when a facility has several forms, the best one counts


def kobo_settings():
    path = os.path.join(ROOT, "private", "kobo.json")
    if os.path.exists(path):
        with open(path, encoding="utf-8") as fh:
            return json.load(fh)
    if os.environ.get("KOBO_TOKEN"):
        return {"server": os.environ.get("KOBO_SERVER", "https://kf.kobotoolbox.org"),
                "form": os.environ["KOBO_FORM"], "token": os.environ["KOBO_TOKEN"]}
    sys.exit("No KoBo settings. Make private/kobo.json first.")


def fetch(k):
    url = f"{k['server'].rstrip('/')}/api/v2/assets/{k['form']}/data/?format=json"
    rows = []
    while url:
        req = urllib.request.Request(url, headers={"Authorization": f"Token {k['token']}"})
        with urllib.request.urlopen(req, timeout=120) as resp:
            page = json.load(resp)
        rows += page["results"]
        url = page.get("next")  # KoBo sends big lists in pages
    return rows


def metres(lat1, lon1, lat2, lon2):
    p = math.pi / 180
    a = (math.sin((lat2 - lat1) * p / 2) ** 2
         + math.cos(lat1 * p) * math.cos(lat2 * p) * math.sin((lon2 - lon1) * p / 2) ** 2)
    return 2 * 6371000 * math.asin(math.sqrt(a))


def judge(sub, fac):
    """Turn one form into (status, distance in m or None, note)."""
    dist, notes = None, []
    gps = str(sub.get("gps") or "").split()
    if len(gps) >= 2 and fac and fac["lat"] is not None:
        dist = round(metres(float(gps[0]), float(gps[1]), fac["lat"], fac["lon"]))
    if len(gps) < 2:
        notes.append("no location recorded")
    elif fac is None or fac["lat"] is None:
        notes.append("facility location not on file yet")
    elif dist > LIMIT_M:
        notes.append(f"{dist:,} m from facility")
    if sub.get("reason"):
        notes.append("reason given (see KoBo)")

    if sub.get("code"):
        status = CONFIRMED  # the form only accepts the right code
    elif dist is not None and dist <= LIMIT_M:
        status = MATCH
    else:
        status = NO_MATCH
    if not sub.get("code"):
        notes.insert(0, "in-charge not available, no code")
    return status, dist, "; ".join(notes)


COLOURS = {CONFIRMED: "E6F4EA", MATCH: "FFF4D6", NO_MATCH: "FDE8E7", NOT_YET: "EEF0F3"}


def write_excel(path, round_name, lines, counts):
    """The same list as status.csv, as an Excel file people can download from the website."""
    wb = openpyxl.Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Summary"
    ws.append([f"{round_name} delivery status"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Updated {datetime.now(LAGOS).strftime('%d %b %Y, %H:%M')}"])
    ws.append([])
    for status, n in counts.items():
        ws.append([status, n])
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=COLOURS[status])
    ws.append([])
    ws.append(["LGA", "Facilities", "Confirmed", "Reported, matches", "Reported, doesn't match", "Not reported yet"])
    for c in ws[ws.max_row]:
        c.font = bold
    by = {}
    for ln in lines:
        x = by.setdefault(ln[1], {CONFIRMED: 0, MATCH: 0, NO_MATCH: 0, NOT_YET: 0})
        x[ln[3]] += 1
    for lga in sorted(by):
        x = by[lga]
        ws.append([lga, sum(x.values()), x[CONFIRMED], x[MATCH], x[NO_MATCH], x[NOT_YET]])
    ws.column_dimensions["A"].width = 42
    for col in "BCDEF":
        ws.column_dimensions[col].width = 16

    ws = wb.create_sheet("Facilities")
    ws.append(["S/N", "LGA", "Facility", "Status", "Date", "Distance (m)", "Note"])
    for c in ws[1]:
        c.font = bold
    order = {NO_MATCH: 0, NOT_YET: 1, MATCH: 2, CONFIRMED: 3}  # problems first
    for ln in sorted(lines, key=lambda x: (order[x[3]], x[1], x[2])):
        ws.append(ln)
        ws.cell(ws.max_row, 4).fill = PatternFill("solid", fgColor=COLOURS[ln[3]])
    for col, width in zip("ABCDEFG", (6, 18, 48, 40, 12, 13, 60)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main():
    with open(os.path.join(DATA, "round.json"), encoding="utf-8") as fh:
        round_name = json.load(fh)["round"]
    with open(os.path.join(DATA, "plan.csv"), encoding="utf-8-sig", newline="") as fh:
        plan = list(csv.DictReader(fh))
    where = {f["sn"]: f for f in load_facilities()}

    subs = fetch(kobo_settings())
    folder = os.path.join(ROOT, "private", round_name)
    os.makedirs(folder, exist_ok=True)
    with open(os.path.join(folder, "submissions.json"), "w", encoding="utf-8") as fh:
        json.dump(subs, fh, ensure_ascii=False, indent=1)

    best = {}  # S/N -> (status, distance, note, date)
    used = 0
    for s in subs:
        if s.get("round") != round_name or not s.get("facility"):
            continue  # a form from another round
        used += 1
        sn = str(s["facility"])
        status, dist, note = judge(s, where.get(sn))
        date = str(s.get("today") or s.get("_submission_time", ""))[:10]
        old = best.get(sn)
        if old is None or (RANK[status], date) > (RANK[old[0]], old[3]):
            best[sn] = (status, dist, note, date)

    counts = {CONFIRMED: 0, MATCH: 0, NO_MATCH: 0, NOT_YET: 0}
    lines = []
    for r in sorted(plan, key=lambda x: (x["LGA"], x["Facility Name"])):
        status, dist, note, date = best.get(r["S/N"], (NOT_YET, None, "", ""))
        counts[status] += 1
        lines.append([r["S/N"], r["LGA"], r["Facility Name"], status, date, "" if dist is None else dist, note])
    with open(os.path.join(DATA, "status.csv"), "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["S/N", "LGA", "Facility", "Status", "Date", "Distance (m)", "Note"])
        w.writerows(lines)
    write_excel(os.path.join(DATA, "status.xlsx"), round_name, lines, counts)
    with open(os.path.join(DATA, "updated.json"), "w", encoding="utf-8") as fh:
        json.dump({"updated": datetime.now(timezone.utc).isoformat(timespec="seconds")}, fh)

    print(f"{round_name}: {len(subs)} forms in KoBo, {used} for this round.")
    for status, n in counts.items():
        print(f"  {n:4}  {status}")
    print("Saved data/status.csv, data/status.xlsx and data/updated.json")


if __name__ == "__main__":
    main()
