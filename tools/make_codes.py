"""Make a new secret 6-digit code for every facility in this round's plan.

Usage (from the project folder, after new_round.py):
    py tools/make_codes.py

Writes, inside private/<round>/ (never sent to GitHub):
    codes - <LGA>.xlsx     one file per LGA, for that LGA's FP coordinator
    code_checks.csv        the scrambled codes, used later to build the KoBo form
"""
import csv
import hashlib
import json
import os
import secrets
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side

from common import DATA, ROOT


def scramble(salt, sn, code):
    """One-way scramble (SHA-256). The KoBo form will do exactly the same to compare."""
    return hashlib.sha256(f"{salt}{sn}{code}".encode()).hexdigest()


def main():
    with open(os.path.join(DATA, "round.json"), encoding="utf-8") as fh:
        round_name = json.load(fh)["round"]
    folder = os.path.join(ROOT, "private", round_name)
    if os.path.exists(os.path.join(folder, "code_checks.csv")):
        sys.exit(f"Codes for {round_name} already exist in {folder}. Delete that folder first to make new ones.")
    os.makedirs(folder, exist_ok=True)

    with open(os.path.join(DATA, "plan.csv"), encoding="utf-8-sig", newline="") as fh:
        plan = list(csv.DictReader(fh))

    salt = secrets.token_hex(8)  # a random word for this round, mixed into every scramble
    by_lga, checks = {}, []
    for row in plan:
        code = f"{secrets.randbelow(1_000_000):06d}"  # random 6 digits, e.g. 048213
        by_lga.setdefault(row["LGA"], []).append((row["Facility Name"], code))
        checks.append((row["S/N"], scramble(salt, row["S/N"], code)))

    with open(os.path.join(folder, "code_checks.csv"), "w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["S/N", "check", "salt"])
        for sn, check in checks:
            writer.writerow([sn, check, salt])

    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    headings = ["Facility", "Code"]

    for lga, rows in by_lga.items():
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Codes"
        ws["A1"] = f"{round_name}: delivery codes for {lga} LGA"
        ws["A1"].font = Font(bold=True, size=16)
        ws["A2"] = ("Give each code ONLY to that facility's in-charge personnel, privately. The in-charge "
                    "personnel tells the code to the delivery person only after the commodities have arrived.")
        ws["A2"].font = Font(size=13)
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells("A2:B2")
        ws.row_dimensions[2].height = 70

        for col, text in enumerate(headings, start=1):
            cell = ws.cell(row=4, column=col, value=text)
            cell.font = Font(bold=True, size=14)
            cell.border = box

        row = 5
        for name, code in sorted(rows):
            values = [name, code]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.font = Font(size=14) if col != 2 else Font(name="Consolas", size=16, bold=True)
                cell.border = box
                cell.alignment = Alignment(vertical="center", wrap_text=(col == 1))
            ws.row_dimensions[row].height = 30
            row += 2  # leave an empty row between facilities

        for col, width in zip("AB", (55, 16)):
            ws.column_dimensions[col].width = width

        # printing: fit the page width, repeat the headings on every page
        ws.print_title_rows = "4:4"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

        wb.save(os.path.join(folder, f"codes - {lga}.xlsx"))

    print(f"{round_name}: made codes for {len(checks)} facilities in {len(by_lga)} LGAs.")
    print(f"Files are in {folder}")


if __name__ == "__main__":
    main()
