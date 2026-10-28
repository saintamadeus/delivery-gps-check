"""Build this round's KoBo form (an XLSForm Excel file).

Usage (from the project folder, after make_codes.py):
    py tools/make_form.py

Writes private/<round>/delivery_form.xlsx  (never sent to GitHub).
Upload it to KoBo: open the form, click "Replace form", pick the file, then DEPLOY.
"""
import csv
import json
import os
import re
from datetime import datetime

import openpyxl

from common import DATA, ROOT, load_facilities


def main():
    with open(os.path.join(DATA, "round.json"), encoding="utf-8") as fh:
        round_name = json.load(fh)["round"]
    folder = os.path.join(ROOT, "private", round_name)

    with open(os.path.join(DATA, "plan.csv"), encoding="utf-8-sig", newline="") as fh:
        plan = list(csv.DictReader(fh))
    with open(os.path.join(folder, "code_checks.csv"), encoding="utf-8", newline="") as fh:
        checks = {r["S/N"]: r for r in csv.DictReader(fh)}
    where = {f["sn"]: f for f in load_facilities()}

    def lga_id(lga):  # "Amuwo Odofin" -> "amuwo_odofin" (KoBo choice names cannot have spaces)
        return re.sub(r"[^a-z0-9]+", "_", lga.lower()).strip("_")

    wb = openpyxl.Workbook()

    # --- the questions ---------------------------------------------------------
    survey = wb.active
    survey.title = "survey"
    survey.append(["type", "name", "label", "required", "appearance", "choice_filter",
                   "constraint", "constraint_message", "relevant", "calculation", "trigger"])
    # distance in metres between the phone and the facility (blank if either location is missing)
    DIST = ("if(${flat} != '' and ${glat} != '', round(6371000 * 2 * asin(sqrt("
            "pow(sin((${glat} - ${flat}) * pi() div 360), 2) + "
            "cos(${flat} * pi() div 180) * cos(${glat} * pi() div 180) * "
            "pow(sin((${glon} - ${flon}) * pi() div 360), 2))), 0), '')")
    HANDED_OVER = "(${code} != '' or ${received_by} != '')"
    NO_GPS = HANDED_OVER + " and ${gps} = ''"
    NEED_REASON = "(${distance_m} != '' and ${distance_m} > 500) or (" + NO_GPS + ")"
    NO_INCHARGE = "selected(${no_incharge}, 'yes')"
    check_now = "digest(concat(${salt}, ${facility}, ${code}), 'SHA-256', 'hex') = ${check}"
    rows = [
        ["start", "start"],
        ["end", "end"],
        ["today", "today"],
        ["calculate", "round", "", "", "", "", "", "", "", f"'{round_name}'"],
        ["text", "deliverer", "Your name (the person delivering)", "yes"],
        ["select_one lga", "lga", "LGA", "yes", "minimal"],
        ["select_one facility", "facility", "Facility", "yes", "minimal", "lga = ${lga}"],
        ["calculate", "check", "", "", "", "", "", "", "",
         "instance('facility')/root/item[name = ${facility}]/check"],
        ["calculate", "salt", "", "", "", "", "", "", "",
         "instance('facility')/root/item[name = ${facility}]/salt"],
        ["select_multiple tick", "no_incharge", "Tick ONLY if the in-charge is not there to give the code"],
        ["text", "code", "Facility code (ask the in-charge AFTER the goods have arrived)", "yes", "numbers", "",
         "digest(concat(${salt}, ${facility}, .), 'SHA-256', 'hex') = ${check}",
         "Wrong code. Ask the facility in-charge again.", "not(" + NO_INCHARGE + ")"],
        ["note", "code_ok", "✅ Code correct. Delivery confirmed by the facility.", "", "", "", "", "", check_now],
        ["text", "received_by", "Who received the goods instead? (name and role, e.g. 'Mrs Ade, nurse on duty')",
         "yes", "", "", "", "", NO_INCHARGE],
        # location: taken by the phone itself the moment the code (or the receiver's name) is typed.
        # Nothing to type or drag.
        ["background-geopoint", "gps_code", "", "", "", "", "", "", "", "", "${code}"],
        ["background-geopoint", "gps_nocode", "", "", "", "", "", "", "", "", "${received_by}"],
        ["calculate", "gps", "", "", "", "", "", "", "", "coalesce(${gps_code}, ${gps_nocode})"],
        ["calculate", "flat", "", "", "", "", "", "", "", "instance('facility')/root/item[name = ${facility}]/lat"],
        ["calculate", "flon", "", "", "", "", "", "", "", "instance('facility')/root/item[name = ${facility}]/lon"],
        ["calculate", "glat", "", "", "", "", "", "", "", "if(${gps} != '', selected-at(${gps}, 0), '')"],
        ["calculate", "glon", "", "", "", "", "", "", "", "if(${gps} != '', selected-at(${gps}, 1), '')"],
        ["calculate", "distance_m", "", "", "", "", "", "", "", DIST],
        ["note", "far", "⚠️ You are about ${distance_m} m from this facility. If you are not at the facility, say why below.",
         "", "", "", "", "", "${distance_m} != '' and ${distance_m} > 500"],
        ["note", "no_gps", "⚠️ Your location could not be recorded. Allow location on this phone, or say why below.",
         "", "", "", "", "", NO_GPS],
        ["text", "reason", "Why are you not at the facility?", "yes", "multiline", "", "", "", NEED_REASON],
    ]
    for r in rows:
        survey.append(r)

    # --- the lists -------------------------------------------------------------
    choices = wb.create_sheet("choices")
    choices.append(["list_name", "name", "label", "lga", "check", "salt", "lat", "lon"])
    choices.append(["tick", "yes", "In-charge not available, no code"])
    for lga in sorted({r["LGA"] for r in plan}):
        choices.append(["lga", lga_id(lga), lga])
    missing = 0
    for r in sorted(plan, key=lambda x: (x["LGA"], x["Facility Name"])):
        c = checks.get(r["S/N"])
        if not c:
            missing += 1
            continue
        choices.append(["facility", r["S/N"], r["Facility Name"], lga_id(r["LGA"]), c["check"], c["salt"],
                        where.get(r["S/N"], {}).get("lat") or "", where.get(r["S/N"], {}).get("lon") or ""])

    # --- form name -------------------------------------------------------------
    settings = wb.create_sheet("settings")
    settings.append(["form_title", "form_id", "version"])
    settings.append([f"LMD delivery check - {round_name}", "lmd_delivery", datetime.now().strftime("%Y%m%d%H%M")])

    out = os.path.join(folder, "delivery_form.xlsx")
    wb.save(out)
    n = len(plan) - missing
    print(f"{round_name}: form made with {n} facilities.")
    if missing:
        print(f"Warning: {missing} facilities have no code. Run make_codes.py again.")
    print(f"File: {out}")


if __name__ == "__main__":
    main()
